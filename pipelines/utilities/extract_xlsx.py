import openpyxl
wb = openpyxl.load_workbook('opus_v_complete_reconstruction_system.xlsx', data_only=True)
print('Sheet names:', wb.sheetnames)
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print('\n' + '='*80)
    print(f'SHEET: {sheet_name}')
    print(f'Rows: {ws.max_row}, Cols: {ws.max_column}')
    print('='*80)
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=False):
        vals = [str(cell.value) if cell.value is not None else '' for cell in row]
        print('\t'.join(vals))
