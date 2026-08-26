from __future__ import annotations

import json
import math
import re
from dataclasses import asdict, dataclass, field, is_dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

FORMULA_SHEETS = ("Brand Reconstructions", "Iris Editions", "Opus V Originals")
DEFAULT_EDP_CONCENTRATION_PCT = 20.0
DEFAULT_FINAL_PRODUCT_DENSITY_G_PER_ML = 1.0
DEFAULT_ETHANOL_DENSITY_G_PER_ML = 0.816
DEFAULT_INGREDIENT_DENSITY_G_PER_ML = 1.0

REPO_BASELINE_PATH = Path("all_formula_data.txt")


@dataclass
class WorkbookMeta:
    path: str
    sheet_names: list[str]
    claimed_formula_counts: dict[str, int]
    actual_formula_counts: dict[str, int]
    ingredient_master_count: int
    substitution_rule_count: int
    total_formula_count: int
    generated_on: str | None = None
    excel_formula_count: int = 0
    data_validation_count: int = 0
    defined_name_count: int = 0
    static_only: bool = True
    overview_notes: list[str] = field(default_factory=list)


@dataclass
class MoleculeRecord:
    code: str
    full_name: str
    cas: str | None = None
    molecular_weight: float | None = None
    logp: float | None = None
    vp_25c_mmhg: float | None = None
    layer: str | None = None
    receptor_pathway: str | None = None
    odor_description: str | None = None
    ifra_limit: str | None = None
    depot_class: str | None = None
    anosmia_risk: str | None = None


@dataclass
class FormulaRow:
    order: int
    layer: str
    code: str
    molecule_name: str
    parts: float
    edp_pct: float
    vp_25c_mmhg: float | None = None
    receptor_pathway: str | None = None
    role_note: str | None = None
    ifra_limit: str | None = None
    style_tags: list[str] = field(default_factory=list)


@dataclass
class FormulaRecord:
    source_sheet: str
    section: str
    name: str
    family: str | None
    base_assignment: str | None
    notes: str | None
    rows: list[FormulaRow]
    total_parts: float
    total_edp_pct: float
    dilution_note: str | None
    style_tags: list[str] = field(default_factory=list)
    structure_mode: str = "formula"

    def ingredients_pct(self) -> dict[str, float]:
        if self.total_parts <= 0:
            return {}
        return {
            row.molecule_name: round((row.parts / self.total_parts) * 100.0, 6)
            for row in self.rows
        }

    def dilutions(self) -> dict[str, float]:
        return {
            row.molecule_name: infer_dilution_fraction(row.molecule_name)
            for row in self.rows
        }

    def verification_payload(
        self,
        batch_volume_ml: float = 30.0,
        edp_concentration_pct: float = DEFAULT_EDP_CONCENTRATION_PCT,
    ) -> dict[str, Any]:
        concentrate_ml = batch_volume_ml * (edp_concentration_pct / 100.0)
        ingredients_pct = self.ingredients_pct()
        ingredients_ul = {
            name: round(concentrate_ml * 1000.0 * (pct / 100.0), 2)
            for name, pct in ingredients_pct.items()
        }
        return {
            "name": self.name,
            "body": self.notes or self.name,
            "structure_mode": self.structure_mode,
            "ingredients_pct": ingredients_pct,
            "ingredients_ul": ingredients_ul,
            "dilutions": self.dilutions(),
            "batch_volume_ml": batch_volume_ml,
            "concentrate_ml": concentrate_ml,
        }


@dataclass
class SubstitutionRule:
    missing_molecule: str
    reason_unavailable: str | None
    best_substitute: str | None
    dose_ratio: str | None
    what_you_lose: str | None
    acceptable: str | None


@dataclass
class WorkbookAudit:
    internal_math_passes: bool
    static_only: bool
    claimed_vs_actual_counts: dict[str, dict[str, int]]
    parts_total_failures: list[str] = field(default_factory=list)
    edp_total_failures: list[str] = field(default_factory=list)
    code_name_failures: list[str] = field(default_factory=list)
    duplicate_formula_names: list[str] = field(default_factory=list)
    missing_or_truncated_sections: list[str] = field(default_factory=list)
    issues: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    passes: list[str] = field(default_factory=list)
    optimization_claim_verdict: str = "unproven"


@dataclass
class WorkbookCatalog:
    meta: WorkbookMeta
    molecules: list[MoleculeRecord]
    formulas: list[FormulaRecord]
    substitutions: list[SubstitutionRule]
    audit: WorkbookAudit | None = None

    def formula_by_name(self, name: str) -> FormulaRecord | None:
        wanted = normalize_compare_name(name)
        for formula in self.formulas:
            if normalize_compare_name(formula.name) == wanted:
                return formula
        return None


def parse_workbook_catalog(path: str | Path) -> WorkbookCatalog:
    workbook_path = Path(path)
    wb = load_workbook(workbook_path, data_only=False, read_only=False)

    claimed_counts, generated_on, overview_notes = _parse_overview_claims(wb["Overview"])
    molecules = _parse_ingredient_master(wb["Ingredient Master"])
    substitutions = _parse_substitution_rules(wb["Fallback & Subs"])
    formulas: list[FormulaRecord] = []
    for sheet_name in FORMULA_SHEETS:
        formulas.extend(_parse_formula_sheet(wb[sheet_name]))

    actual_counts = {
        sheet_name: sum(1 for formula in formulas if formula.source_sheet == sheet_name)
        for sheet_name in FORMULA_SHEETS
    }
    excel_formula_count = _count_excel_formulas(wb)
    data_validation_count = sum(len(ws.data_validations.dataValidation) for ws in wb.worksheets)
    defined_name_count = len(list(wb.defined_names))

    meta = WorkbookMeta(
        path=str(workbook_path.resolve()),
        sheet_names=list(wb.sheetnames),
        claimed_formula_counts=claimed_counts,
        actual_formula_counts=actual_counts,
        ingredient_master_count=len(molecules),
        substitution_rule_count=len(substitutions),
        total_formula_count=len(formulas),
        generated_on=generated_on,
        excel_formula_count=excel_formula_count,
        data_validation_count=data_validation_count,
        defined_name_count=defined_name_count,
        static_only=(excel_formula_count == 0 and data_validation_count == 0 and defined_name_count == 0),
        overview_notes=overview_notes,
    )
    catalog = WorkbookCatalog(meta=meta, molecules=molecules, formulas=formulas, substitutions=substitutions)
    catalog.audit = audit_workbook_catalog(catalog)
    return catalog


def audit_workbook_catalog(catalog: WorkbookCatalog) -> WorkbookAudit:
    molecule_by_code = {normalize_key(m.code): m for m in catalog.molecules}
    parts_failures: list[str] = []
    edp_failures: list[str] = []
    code_name_failures: list[str] = []
    missing_sections: list[str] = []
    seen: dict[str, int] = {}

    for formula in catalog.formulas:
        row_parts_total = round(sum(row.parts for row in formula.rows), 6)
        row_edp_total = round(sum(row.edp_pct for row in formula.rows), 6)
        if not math.isclose(row_parts_total, formula.total_parts, rel_tol=0.0, abs_tol=0.05):
            parts_failures.append(
                f"{formula.name}: rows={row_parts_total:.4f} total={formula.total_parts:.4f}"
            )
        if not math.isclose(row_edp_total, formula.total_edp_pct, rel_tol=0.0, abs_tol=0.15):
            edp_failures.append(
                f"{formula.name}: rows={row_edp_total:.4f} total={formula.total_edp_pct:.4f}"
            )
        if not formula.rows or formula.total_parts <= 0:
            missing_sections.append(f"{formula.name}: missing formula rows or total")
        seen[normalize_key(formula.name)] = seen.get(normalize_key(formula.name), 0) + 1
        for row in formula.rows:
            molecule = molecule_by_code.get(normalize_key(row.code))
            if molecule is None:
                code_name_failures.append(f"{formula.name}: unknown code {row.code} for {row.molecule_name}")
                continue
            if normalize_compare_name(molecule.full_name) != normalize_compare_name(row.molecule_name):
                code_name_failures.append(
                    f"{formula.name}: {row.code} workbook='{row.molecule_name}' master='{molecule.full_name}'"
                )

    claimed_vs_actual = {
        sheet_name: {
            "claimed": catalog.meta.claimed_formula_counts.get(sheet_name, 0),
            "actual": catalog.meta.actual_formula_counts.get(sheet_name, 0),
        }
        for sheet_name in FORMULA_SHEETS
    }
    duplicate_names = sorted(
        formula_name for formula_name, count in seen.items() if count > 1
    )

    issues: list[str] = []
    warnings: list[str] = []
    passes: list[str] = []

    count_mismatches = [
        f"{sheet}: claimed {counts['claimed']} vs actual {counts['actual']}"
        for sheet, counts in claimed_vs_actual.items()
        if counts["claimed"] and counts["claimed"] != counts["actual"]
    ]
    if count_mismatches:
        issues.extend(count_mismatches)
    else:
        passes.append("Claimed sheet counts match parsed counts.")

    if not parts_failures and not edp_failures:
        passes.append("All formula blocks reconcile Parts totals and EDP totals.")
    if not code_name_failures:
        passes.append("All formula ingredient codes resolve cleanly against the ingredient master.")
    if catalog.meta.static_only:
        passes.append("Workbook is static-only: no Excel formulas, data validations, or computational artifacts were found.")

    warnings.append(
        "Workbook uses 'optimized' and 'IFRA-compliant' as authored labels; no in-file pre/post optimization history or scoring trace exists."
    )

    internal_math_passes = not parts_failures and not edp_failures and not code_name_failures
    return WorkbookAudit(
        internal_math_passes=internal_math_passes,
        static_only=catalog.meta.static_only,
        claimed_vs_actual_counts=claimed_vs_actual,
        parts_total_failures=parts_failures,
        edp_total_failures=edp_failures,
        code_name_failures=code_name_failures,
        duplicate_formula_names=duplicate_names,
        missing_or_truncated_sections=missing_sections,
        issues=issues,
        warnings=warnings,
        passes=passes,
        optimization_claim_verdict="unproven",
    )


def run_engine_validation(record: FormulaRecord) -> dict[str, Any]:
    result: dict[str, Any] = {
        "formula_name": record.name,
        "validation_status": "not_run",
        "rule_coverage": None,
        "score_summary": None,
        "errors": [],
    }
    formula_pct = record.ingredients_pct()

    try:
        from engine.optimizer.models import FormulaVector, analyze_formula_rule_coverage
        from engine.optimizer.scoring import FormulaScorer
        from engine.validator import validate_formula
    except Exception as exc:
        result["errors"].append(f"Engine import failed: {exc}")
        return result

    try:
        validation = validate_formula(formula_pct, product_type="fine_fragrance")
        result["validation_status"] = "pass" if getattr(validation, "is_valid", False) else "fail"
        result["validator"] = {
            "is_valid": bool(getattr(validation, "is_valid", False)),
            "overall_score": getattr(validation, "overall_score", None),
            "errors": list(getattr(validation, "errors", []) or []),
            "warnings": list(getattr(validation, "warnings", []) or []),
        }
    except Exception as exc:
        result["errors"].append(f"Validator failed: {exc}")

    try:
        rule_coverage = analyze_formula_rule_coverage(list(formula_pct.keys()))
        result["rule_coverage"] = {
            "positive_pairs": sorted(list(rule_coverage.get("positive_pairs", set()))),
            "conflict_pairs": sorted(list(rule_coverage.get("conflict_pairs", set()))),
            "covered_pairs": sorted(list(rule_coverage.get("covered_pairs", set()))),
            "total_pairs": rule_coverage.get("total_pairs"),
        }
    except Exception as exc:
        result["errors"].append(f"Rule coverage failed: {exc}")

    try:
        vector = FormulaVector(ingredients=formula_pct, dilutions=record.dilutions())
        scorer = FormulaScorer()
        score = scorer.score(vector)
        result["score_summary"] = {
            "overall": score.get("overall"),
            "style_match": score.get("style_match"),
            "identity": score.get("identity"),
            "balance": score.get("balance"),
            "texture": score.get("texture"),
            "tenacity": score.get("longevity"),
        }
    except Exception as exc:
        result["errors"].append(f"Scoring failed: {exc}")

    return result


def build_verification_bundle_for_formula(
    record: FormulaRecord,
    batch_volume_ml: float = 30.0,
    edp_concentration_pct: float = DEFAULT_EDP_CONCENTRATION_PCT,
) -> dict[str, Any]:
    payload = record.verification_payload(
        batch_volume_ml=batch_volume_ml,
        edp_concentration_pct=edp_concentration_pct,
    )
    from scripts.verify_formula_workflow import build_verification_bundle

    return build_verification_bundle(payload)


def build_inventory_crosswalk(
    record: FormulaRecord,
    substitutions: list[SubstitutionRule],
) -> dict[str, Any]:
    from engine.inventory_parser import parse_inventory

    inventory = parse_inventory(include_solvents=True, include_unavailable=False)
    inventory_index: dict[str, list[str]] = {}
    for material in inventory:
        candidates = {
            normalize_inventory_name(material.name),
            normalize_inventory_name(getattr(material, "raw_name", "") or ""),
        }
        for candidate in candidates:
            if candidate:
                inventory_index.setdefault(candidate, []).append(material.name)

    substitution_index = {
        normalize_inventory_name(rule.missing_molecule): rule for rule in substitutions
    }

    rows: list[dict[str, Any]] = []
    direct_available = 0
    substituted = 0
    missing = 0
    for row in record.rows:
        normalized = normalize_inventory_name(row.molecule_name)
        direct_match = inventory_index.get(normalized, [])
        if direct_match:
            direct_available += 1
            rows.append(
                {
                    "ingredient": row.molecule_name,
                    "code": row.code,
                    "status": "available",
                    "inventory_match": direct_match[0],
                    "substitution": None,
                }
            )
            continue

        substitute_rule = substitution_index.get(normalized)
        substitute_match = None
        if substitute_rule and substitute_rule.best_substitute:
            sub_normalized = normalize_inventory_name(substitute_rule.best_substitute)
            matches = inventory_index.get(sub_normalized, [])
            if matches:
                substitute_match = matches[0]

        if substitute_rule and substitute_match:
            substituted += 1
            rows.append(
                {
                    "ingredient": row.molecule_name,
                    "code": row.code,
                    "status": "substituted",
                    "inventory_match": None,
                    "substitution": {
                        "best_substitute": substitute_rule.best_substitute,
                        "inventory_match": substitute_match,
                        "dose_ratio": substitute_rule.dose_ratio,
                        "acceptable": substitute_rule.acceptable,
                    },
                }
            )
        else:
            missing += 1
            rows.append(
                {
                    "ingredient": row.molecule_name,
                    "code": row.code,
                    "status": "missing",
                    "inventory_match": None,
                    "substitution": asdict(substitute_rule) if substitute_rule else None,
                }
            )

    return {
        "formula_name": record.name,
        "available_count": direct_available,
        "substituted_count": substituted,
        "missing_count": missing,
        "rows": rows,
    }


def export_formula_weight_first(
    record: FormulaRecord,
    target_ml: float,
    edp_concentration_pct: float = DEFAULT_EDP_CONCENTRATION_PCT,
    final_product_density_g_per_ml: float = DEFAULT_FINAL_PRODUCT_DENSITY_G_PER_ML,
    ethanol_density_g_per_ml: float = DEFAULT_ETHANOL_DENSITY_G_PER_ML,
    ingredient_density_g_per_ml: float = DEFAULT_INGREDIENT_DENSITY_G_PER_ML,
    variant_name: str = "workbook_original",
) -> list[dict[str, Any]]:
    final_mass_g = target_ml * final_product_density_g_per_ml
    concentrate_mass_g = final_mass_g * (edp_concentration_pct / 100.0)
    ethanol_mass_g = final_mass_g - concentrate_mass_g
    export_rows: list[dict[str, Any]] = []
    for row in record.rows:
        ingredient_mass_g = concentrate_mass_g * (row.parts / record.total_parts)
        estimated_ml = ingredient_mass_g / ingredient_density_g_per_ml if ingredient_density_g_per_ml else None
        export_rows.append(
            {
                "formula_name": record.name,
                "variant": variant_name,
                "target_ml": target_ml,
                "phase": row.layer,
                "ingredient": row.molecule_name,
                "code": row.code,
                "parts": row.parts,
                "concentrate_pct": round((row.parts / record.total_parts) * 100.0, 6),
                "ingredient_g": round(ingredient_mass_g, 6),
                "estimated_ml": round(estimated_ml, 6) if estimated_ml is not None else None,
                "estimated_ul": round(estimated_ml * 1000.0, 2) if estimated_ml is not None else None,
                "note": "Approximate volume estimate only; mass is canonical.",
            }
        )

    ethanol_ml = ethanol_mass_g / ethanol_density_g_per_ml if ethanol_density_g_per_ml else None
    export_rows.append(
        {
            "formula_name": record.name,
            "variant": variant_name,
            "target_ml": target_ml,
            "phase": "SOLVENT",
            "ingredient": "95% ethanol",
            "code": "ETOH95",
            "parts": None,
            "concentrate_pct": None,
            "ingredient_g": round(ethanol_mass_g, 6),
            "estimated_ml": round(ethanol_ml, 6) if ethanol_ml is not None else None,
            "estimated_ul": round(ethanol_ml * 1000.0, 2) if ethanol_ml is not None else None,
            "note": "Approximate ethanol volume estimate from density; mass is canonical.",
        }
    )
    return export_rows


def compare_formula_against_baselines(
    record: FormulaRecord,
    baseline_path: str | Path = REPO_BASELINE_PATH,
) -> dict[str, Any]:
    baselines = _load_repo_baselines(Path(baseline_path))
    target = {
        normalize_compare_name(row.molecule_name): {
            "parts": row.parts,
            "layer": row.layer,
            "receptor": row.receptor_pathway,
        }
        for row in record.rows
    }
    target_total = sum(item["parts"] for item in target.values()) or 1.0

    repo_reports: list[dict[str, Any]] = []
    for baseline in baselines:
        shared = sorted(set(target) & set(baseline["materials"]))
        added = sorted(set(target) - set(baseline["materials"]))
        removed = sorted(set(baseline["materials"]) - set(target))
        overlap_target_weight = sum(min(target[name]["parts"], baseline["materials"][name]["parts"]) for name in shared)
        abs_diff = sum(
            abs(target.get(name, {"parts": 0.0})["parts"] - baseline["materials"].get(name, {"parts": 0.0})["parts"])
            for name in sorted(set(target) | set(baseline["materials"]))
        )

        target_layer = _layer_distribution_from_rows(record.rows)
        baseline_layer = _layer_distribution_from_materials(baseline["materials"])
        receptor_overlap = _receptor_overlap(record.rows, baseline["materials"])
        sub_odt_shift = _sub_odt_shift(record.rows, baseline["materials"])

        repo_reports.append(
            {
                "baseline_family": baseline["family"],
                "baseline_variant": baseline["variant"],
                "shared_ingredients": shared,
                "added_ingredients": added,
                "removed_ingredients": removed,
                "weighted_overlap_pct_of_target": round((overlap_target_weight / target_total) * 100.0, 2),
                "weighted_absolute_drift_pct_of_target": round((abs_diff / target_total) * 100.0, 2),
                "layer_drift": {
                    layer: round(target_layer.get(layer, 0.0) - baseline_layer.get(layer, 0.0), 4)
                    for layer in sorted(set(target_layer) | set(baseline_layer))
                },
                "receptor_overlap_pct": round(receptor_overlap * 100.0, 2),
                "sub_odt_load_shift_pct_points": round(sub_odt_shift, 2),
                "change_classification": classify_drift(overlap_target_weight / target_total, abs_diff / target_total, added, removed),
            }
        )

    official_alignment = _official_brief_alignment(record)
    recommended = _recommend_repo_verdict(repo_reports)
    return {
        "formula_name": record.name,
        "optimization_verdict": ["unproven", recommended],
        "repo_baselines": repo_reports,
        "official_brief_alignment": official_alignment,
    }


def classify_drift(overlap_ratio: float, abs_diff_ratio: float, added: list[str], removed: list[str]) -> str:
    structural_changes = len(added) + len(removed)
    if overlap_ratio >= 0.8 and abs_diff_ratio <= 0.25 and structural_changes <= 3:
        return "close refinement"
    if overlap_ratio >= 0.65 and abs_diff_ratio <= 0.35 and structural_changes <= 5:
        return "shell optimization"
    return "substantial redesign"


def write_json(path: str | Path, payload: Any) -> None:
    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8") as handle:
        json.dump(_json_ready(payload), handle, indent=2, ensure_ascii=False, default=str)


def catalog_to_dict(catalog: WorkbookCatalog) -> dict[str, Any]:
    return asdict(catalog)


def audit_to_markdown(catalog: WorkbookCatalog, drift_report: dict[str, Any] | None = None) -> str:
    audit = catalog.audit
    if audit is None:
        raise ValueError("Catalog audit is missing.")

    lines = [
        "# Opus V Luxury Workbook Audit",
        "",
        f"- Workbook: `{catalog.meta.path}`",
        f"- Sheets: {len(catalog.meta.sheet_names)}",
        f"- Parsed formulas: {catalog.meta.total_formula_count}",
        f"- Ingredient master rows: {catalog.meta.ingredient_master_count}",
        f"- Substitution rules: {catalog.meta.substitution_rule_count}",
        "",
        "## Audit Summary",
        "",
    ]

    for item in audit.passes:
        lines.append(f"- PASS: {item}")
    for item in audit.warnings:
        lines.append(f"- WARN: {item}")
    for item in audit.issues:
        lines.append(f"- ISSUE: {item}")

    if drift_report:
        lines.extend(
            [
                "",
                "## Drift Verdict",
                "",
                f"- Formula: `{drift_report['formula_name']}`",
                f"- Verdict: {', '.join(drift_report['optimization_verdict'])}",
            ]
        )
        for repo_baseline in drift_report.get("repo_baselines", []):
            lines.append(
                "- "
                f"{repo_baseline['baseline_family']} / {repo_baseline['baseline_variant']}: "
                f"overlap {repo_baseline['weighted_overlap_pct_of_target']}%, "
                f"drift {repo_baseline['weighted_absolute_drift_pct_of_target']}%, "
                f"classification {repo_baseline['change_classification']}"
            )

    return "\n".join(lines) + "\n"


def normalize_key(value: str | None) -> str:
    return re.sub(r"[^a-z0-9]+", "", (value or "").strip().lower())


def normalize_compare_name(value: str | None) -> str:
    text = (value or "").strip().lower()
    replacements = {
        "amyl methyl ionone (fixolide ionone)": "amyl methyl ionone",
        "koavone (4-methyl ionone)": "koavone",
        "ultralia (methyl ionone alpha)": "ultralia",
        "orivone (methyl ionone gamma)": "orivone",
        "ambroxan 30% dpg (ambrofix)": "ambroxan 30%",
        "galaxolide 50% ipm": "galaxolide 50%",
        "habanolide (exaltolide)": "habanolide",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    text = re.sub(r"\([^)]*\)", "", text)
    text = re.sub(r"[^a-z0-9%]+", " ", text)
    return " ".join(text.split())


def normalize_inventory_name(value: str | None) -> str:
    text = (value or "").strip().lower()
    code_prefixed = re.match(r"^([a-z0-9]{2,4})\s*\(([^)]+)\)(?:.*)?$", text)
    if code_prefixed:
        text = code_prefixed.group(2).strip().lower()
    aliases = {
        "d-limonene": "limonene",
        "phenyl ethyl alcohol": "phenethyl alcohol",
        "phenyl ethyl alcohol (pea)": "phenethyl alcohol",
        "myristic acid powder": "myristic acid",
        "ambroxan 30% dpg (ambrofix)": "ambrofix",
        "ambroxan 30% dpg": "ambrofix",
        "galaxolide 50% ipm": "galaxolide",
        "habanolide (exaltolide)": "habanolide",
        "amyl methyl ionone (fixolide ionone)": "amyl methyl ionone",
        "koavone (4-methyl ionone)": "koavone",
        "ultralia (methyl ionone alpha)": "ultralia",
        "orivone (methyl ionone gamma)": "orivone",
        "musk ketone 10% dpg": "musk ketone 10",
        "musk ketone 10% dep": "musk ketone 10",
        "musk ketone (10%)": "musk ketone 10",
        "cyclamal (cyclamen aldehyde)": "cyclamen aldehyde",
    }
    if text in aliases:
        return aliases[text]
    text = re.sub(r"\([^)]*\)", "", text)
    text = text.replace("%", " ")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def infer_dilution_fraction(molecule_name: str) -> float:
    match = re.search(r"(\d+(?:\.\d+)?)\s*%", molecule_name)
    if not match:
        return 1.0
    return round(float(match.group(1)) / 100.0, 6)


def _parse_overview_claims(ws) -> tuple[dict[str, int], str | None, list[str]]:
    claims: dict[str, int] = {}
    generated_on: str | None = None
    notes: list[str] = []
    for row in ws.iter_rows(values_only=True):
        values = [value for value in row if value not in (None, "")]
        if len(values) < 2:
            continue
        key = str(values[0]).strip()
        detail = str(values[1]).strip()
        if key == "Workbook Generated":
            generated_on = detail
        if key in FORMULA_SHEETS:
            count_match = re.search(r"(\d+)", detail)
            if count_match:
                claims[key] = int(count_match.group(1))
        if key in {"HOW TO USE", "Formula Notes", "Color Code"}:
            notes.append(detail)
    return claims, generated_on, notes


def _parse_ingredient_master(ws) -> list[MoleculeRecord]:
    molecules: list[MoleculeRecord] = []
    header_row = None
    for row in ws.iter_rows():
        values = [cell.value for cell in row if cell.value not in (None, "")]
        if values and str(values[0]).strip() == "Code":
            header_row = row
            break
    if header_row is None:
        return molecules

    start_row = header_row[0].row + 1
    for row in ws.iter_rows(min_row=start_row):
        values = [cell.value for cell in row[1:13]]
        code = _as_text(values[0])
        full_name = _as_text(values[1])
        if not code or not full_name:
            continue
        molecules.append(
            MoleculeRecord(
                code=code,
                full_name=full_name,
                cas=_as_text(values[2]),
                molecular_weight=_to_float(values[3]),
                logp=_to_float(values[4]),
                vp_25c_mmhg=_to_float(values[5]),
                layer=_as_text(values[6]),
                receptor_pathway=_as_text(values[7]),
                odor_description=_as_text(values[8]),
                ifra_limit=_as_text(values[9]),
                depot_class=_as_text(values[10]),
                anosmia_risk=_as_text(values[11]),
            )
        )
    return molecules


def _parse_substitution_rules(ws) -> list[SubstitutionRule]:
    substitutions: list[SubstitutionRule] = []
    started = False
    for row in ws.iter_rows(values_only=True):
        values = [value for value in row if value not in (None, "")]
        if not values:
            continue
        first = str(values[0]).strip()
        if first == "Missing Molecule":
            started = True
            continue
        if not started:
            continue
        if first.startswith("D.") or first.startswith("Notes"):
            break
        if len(values) < 2:
            continue
        substitutions.append(
            SubstitutionRule(
                missing_molecule=_as_text(values[0]) or "",
                reason_unavailable=_as_text(values[1]),
                best_substitute=_as_text(values[2]) if len(values) > 2 else None,
                dose_ratio=_as_text(values[3]) if len(values) > 3 else None,
                what_you_lose=_as_text(values[4]) if len(values) > 4 else None,
                acceptable=_as_text(values[5]) if len(values) > 5 else None,
            )
        )
    return substitutions


def _parse_formula_sheet(ws) -> list[FormulaRecord]:
    formulas: list[FormulaRecord] = []
    current_name: str | None = None
    current_family: str | None = None
    current_base: str | None = None
    current_notes: list[str] = []
    current_rows: list[FormulaRow] = []
    current_style_tags: set[str] = set()
    current_section = ws.title
    total_parts = 0.0
    total_edp = 0.0
    dilution_note: str | None = None

    for row in ws.iter_rows():
        non_empty_cells = [cell for cell in row if cell.value not in (None, "")]
        if not non_empty_cells:
            continue
        values = [cell.value for cell in non_empty_cells]
        text = _as_text(values[0]) or ""

        if _is_formula_title(ws.title, values) and (current_name is None or current_rows):
            if current_name and current_rows:
                formulas.append(
                    FormulaRecord(
                        source_sheet=ws.title,
                        section=current_section,
                        name=current_name,
                        family=current_family,
                        base_assignment=current_base,
                        notes=" ".join(current_notes).strip() or None,
                        rows=current_rows[:],
                        total_parts=total_parts,
                        total_edp_pct=total_edp,
                        dilution_note=dilution_note,
                        style_tags=sorted(current_style_tags),
                        structure_mode=_infer_structure_mode(current_name, current_notes),
                    )
                )
            current_name = text
            current_family = None
            current_base = None
            current_notes = []
            current_rows = []
            current_style_tags = set(_extract_style_tags(non_empty_cells))
            current_section = ws.title
            total_parts = 0.0
            total_edp = 0.0
            dilution_note = None
            continue

        if current_name is None:
            continue

        if text.startswith("Family:"):
            current_family, current_base = _parse_family_and_base(text)
            continue

        if text.startswith("Source:") or text.startswith("Theme:"):
            current_notes.append(text)
            continue

        if text.startswith("TOTAL"):
            if len(values) > 1:
                total_parts = _to_float(values[1]) or 0.0
            if len(values) > 2:
                total_edp = _to_float(values[2]) or 0.0
            if len(values) > 3:
                dilution_note = _as_text(values[3])
            formulas.append(
                FormulaRecord(
                    source_sheet=ws.title,
                    section=current_section,
                    name=current_name,
                    family=current_family,
                    base_assignment=current_base,
                    notes=" ".join(current_notes).strip() or None,
                    rows=current_rows[:],
                    total_parts=total_parts,
                    total_edp_pct=total_edp,
                    dilution_note=dilution_note,
                    style_tags=sorted(current_style_tags),
                    structure_mode=_infer_structure_mode(current_name, current_notes),
                )
            )
            current_name = None
            current_family = None
            current_base = None
            current_notes = []
            current_rows = []
            current_style_tags = set()
            total_parts = 0.0
            total_edp = 0.0
            dilution_note = None
            continue

        maybe_order = _to_int(values[0])
        if maybe_order is not None and len(values) >= 6:
            row_style = _extract_style_tags(non_empty_cells)
            current_style_tags.update(row_style)
            current_rows.append(
                FormulaRow(
                    order=maybe_order,
                    layer=_as_text(values[1]) or "",
                    code=_as_text(values[2]) or "",
                    molecule_name=_as_text(values[3]) or "",
                    parts=_to_float(values[4]) or 0.0,
                    edp_pct=_to_float(values[5]) or 0.0,
                    vp_25c_mmhg=_to_float(values[6]) if len(values) > 6 else None,
                    receptor_pathway=_as_text(values[7]) if len(values) > 7 else None,
                    role_note=_as_text(values[8]) if len(values) > 8 else None,
                    ifra_limit=_as_text(values[9]) if len(values) > 9 else None,
                    style_tags=row_style,
                )
            )
            continue

        if not _is_header_or_section_row(text):
            current_notes.append(text)

    return formulas


def _parse_family_and_base(text: str) -> tuple[str | None, str | None]:
    family = None
    base = None
    parts = [part.strip() for part in text.split("|")]
    for part in parts:
        if part.startswith("Family:"):
            family = part.split(":", 1)[1].strip()
        elif part.startswith("Base:"):
            base = part.split(":", 1)[1].strip()
    return family, base


def _is_formula_title(sheet_name: str, values: list[Any]) -> bool:
    if len(values) != 1:
        return False
    text = _as_text(values[0]) or ""
    if not text or _is_header_or_section_row(text):
        return False
    if text.startswith(("Workbook", "Formula", "Sources:", "Method:", "All formulas:")):
        return False
    if text.startswith("Family:") or text.startswith("Source:") or text.startswith("Theme:"):
        return False
    if text.startswith(("⬆", "♦", "▼")):
        return False
    if ":" in text:
        return False
    upper = text.upper()
    if upper.startswith("BRAND RECONSTRUCTIONS") or upper.startswith("OPUS V IRIS EDITIONS") or upper.startswith("OPUS V ORIGINAL CREATIONS"):
        return False
    return sheet_name in FORMULA_SHEETS


def _is_header_or_section_row(text: str) -> bool:
    markers = (
        "#",
        "LAYER",
        "MOLECULE NAME",
        "PARTS",
        "EDP%",
        "TOP NOTES",
        "HEART NOTES",
        "BASE NOTES",
        "BRAND RECONSTRUCTIONS",
        "IRIS EDITIONS",
        "ORIGINAL",
    )
    upper = text.upper()
    return any(marker in upper for marker in markers)


def _infer_structure_mode(name: str, notes: list[str]) -> str:
    joined = " ".join(notes).lower()
    if "accord" in name.lower() or "not a complete perfume" in joined or "building block" in joined:
        return "module"
    return "formula"


def _extract_style_tags(cells: list[Any]) -> list[str]:
    tags: set[str] = set()
    fill_values = {_rgb_value(getattr(cell.fill, "fgColor", None)) for cell in cells}
    font_values = {_rgb_value(getattr(cell.font, "color", None)) for cell in cells}

    if {"00FFF2CC", "00FFF9C4"} & fill_values:
        tags.add("original")
    if "00DDEEFF" in fill_values:
        tags.add("iris_block")
    if "00E2EFDA" in fill_values:
        tags.add("luxury_sub_odt")
    if "00CC0000" in font_values:
        tags.add("ifra_warning")
    return sorted(tag for tag in tags if tag)


def _rgb_value(color: Any) -> str | None:
    if color is None:
        return None
    if getattr(color, "type", None) == "rgb":
        return getattr(color, "rgb", None)
    return None


def _count_excel_formulas(workbook) -> int:
    count = 0
    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    count += 1
    return count


def _load_repo_baselines(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    baselines: dict[tuple[str, str], dict[str, Any]] = {}
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            parts = [item.strip() for item in line.split("|")]
            if len(parts) < 5:
                continue
            family, variant, layer, ingredient, parts_value = parts[:5]
            if family not in {"Opus V Base", "Opus V Luxury"}:
                continue
            key = (family, variant)
            entry = baselines.setdefault(key, {"family": family, "variant": variant, "materials": {}})
            entry["materials"][normalize_compare_name(ingredient)] = {
                "ingredient": ingredient,
                "parts": float(parts_value),
                "layer": layer,
                "receptor": parts[7].strip() if len(parts) > 7 else "",
            }
    return list(baselines.values())


def _official_brief_alignment(record: FormulaRecord) -> dict[str, Any]:
    ingredient_names = {normalize_compare_name(row.molecule_name) for row in record.rows}
    note_tokens = set(" ".join(filter(None, [record.name, record.notes or "", record.family or "", record.base_assignment or ""])).lower().split())

    brief_facets = {
        "iris_or_orris_core": {"alpha irone", "alpha ionone", "beta ionone", "orivone", "ultralia", "koavone", "orris", "iris"},
        "jasmine_floral_lift": {"hedione", "hedione hc", "jasmine", "linalool"},
        "rose_facet": {"rose", "rose oxide", "geraniol", "phenethyl alcohol"},
        "rum_liquor_facet": {"rum", "rhum", "caramel", "ethyl maltol"},
        "oud_dry_woods": {"oud", "agarwood", "ambroxan", "amberwood", "cashmeran", "cedarwood"},
        "animalic_civet": {"civet", "animalic", "castoreum"},
    }

    present: dict[str, list[str]] = {}
    missing: list[str] = []
    for facet, candidates in brief_facets.items():
        hits = sorted(
            candidate
            for candidate in candidates
            if candidate in ingredient_names or candidate in note_tokens
        )
        if hits:
            present[facet] = hits
        else:
            missing.append(facet)

    return {
        "present_facets": present,
        "missing_facets": missing,
        "notes": "Official brief alignment is inferred from repo reverse-engineering context and material vocabulary, not from workbook-native scoring data.",
    }


def _layer_distribution_from_rows(rows: list[FormulaRow]) -> dict[str, float]:
    total = sum(row.parts for row in rows) or 1.0
    result: dict[str, float] = {}
    for row in rows:
        result[row.layer] = result.get(row.layer, 0.0) + (row.parts / total)
    return result


def _layer_distribution_from_materials(materials: dict[str, dict[str, Any]]) -> dict[str, float]:
    total = sum(item["parts"] for item in materials.values()) or 1.0
    result: dict[str, float] = {}
    for item in materials.values():
        layer = item.get("layer") or "UNKNOWN"
        result[layer] = result.get(layer, 0.0) + (item["parts"] / total)
    return result


def _receptor_overlap(rows: list[FormulaRow], baseline_materials: dict[str, dict[str, Any]]) -> float:
    target_receptors = {
        normalize_key(row.receptor_pathway)
        for row in rows
        if row.receptor_pathway
    }
    baseline_receptors = {
        normalize_key(material.get("receptor"))
        for material in baseline_materials.values()
        if material.get("receptor")
    }
    if not target_receptors or not baseline_receptors:
        return 0.0
    return len(target_receptors & baseline_receptors) / len(target_receptors)


def _sub_odt_shift(rows: list[FormulaRow], baseline_materials: dict[str, dict[str, Any]]) -> float:
    target_total = sum(row.parts for row in rows) or 1.0
    baseline_total = sum(item["parts"] for item in baseline_materials.values()) or 1.0
    target_sub_odt = sum(row.parts for row in rows if "sub-odt" in (row.role_note or "").lower())
    baseline_sub_odt = sum(
        item["parts"] for item in baseline_materials.values() if "sub-odt" in str(item.get("ingredient", "")).lower()
    )
    return ((target_sub_odt / target_total) - (baseline_sub_odt / baseline_total)) * 100.0


def _recommend_repo_verdict(repo_reports: list[dict[str, Any]]) -> str:
    if not repo_reports:
        return "unproven"
    closest = max(repo_reports, key=lambda report: report["weighted_overlap_pct_of_target"])
    return closest["change_classification"]


def _to_float(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        text = str(value).strip().replace(",", "")
        if not text:
            return None
        match = re.search(r"-?\d+(?:\.\d+)?", text)
        return float(match.group(0)) if match else None


def _to_int(value: Any) -> int | None:
    try:
        if value in (None, ""):
            return None
        return int(value)
    except (TypeError, ValueError):
        text = str(value).strip()
        return int(text) if text.isdigit() else None


def _as_text(value: Any) -> str | None:
    if value in (None, ""):
        return None
    return str(value).strip()


def _json_ready(value: Any) -> Any:
    if is_dataclass(value):
        return asdict(value)
    if isinstance(value, dict):
        return {k: _json_ready(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_ready(item) for item in value]
    if isinstance(value, set):
        return sorted(_json_ready(item) for item in value)
    return value
