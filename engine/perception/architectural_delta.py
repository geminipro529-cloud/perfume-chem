"""Target-first architectural deltas for complexity experiments.

The engine proposes evidence-producing comparisons.  It never treats material
count as complexity and never grants formula, compounding, sensory, purchase,
safety, or release authority.
"""

from __future__ import annotations

import hashlib
import json
import os
from dataclasses import dataclass, field
from enum import Enum
from pathlib import Path
from typing import Any, Mapping

from openpyxl import load_workbook

from engine.evidence.augmentation import (
    DecisionDeltaV1,
    EvidenceAugmentationState,
    EvidenceDeltaReceiptV1,
    hold_receipt,
    no_augmentation_receipt,
)
from engine.evidence_contracts import canonical_json_bytes, sha256_hex
from engine.perception import complexity_inventory as _complexity_inventory
from engine.perception.complexity_inventory import (
    ComplexityInventoryCatalog,
    InventoryProjection,
)
from engine.scientific_validation.complexity_design_contracts import (
    NaryAssessmentState,
    NaryInteractionCandidate,
    evaluate_nary_interaction,
)


class ArchitecturalDeltaState(str, Enum):
    NO_CHANGE = "NO_CHANGE"
    PROPOSED = "PROPOSED"
    HOLD = "HOLD"


class ArchitecturalDeltaKind(str, Enum):
    OMISSION = "OMISSION"
    ADDITION = "ADDITION"
    RATIO = "RATIO"
    NARY_DESIGN = "NARY_DESIGN"


class ArchitecturalDeltaFamily(str, Enum):
    GENERAL = "GENERAL"
    CITRUS = "CITRUS"
    MUSK = "MUSK"


_CURRENT_MASTER_SHEET = "Current Inventory Master"
_CURRENT_MASTER_HEADERS = {
    "Canonical material": "canonical_material",
    "Status": "status",
    "Actual stock(s)": "actual_stocks",
    "Can prepare": "can_prepare",
    "Family": "family",
    "Alias / non-equivalent": "alias_or_non_equivalent",
    "Formula-use policy": "formula_use_policy",
    "User note": "user_note",
}


def _text(value: object, field_name: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{field_name} must be nonblank text")
    return " ".join(value.split())


def _sha256(value: object, field_name: str) -> str:
    text = _text(value, field_name).lower()
    if len(text) != 64 or any(character not in "0123456789abcdef" for character in text):
        raise ValueError(f"{field_name} must be a SHA-256 hex digest")
    return text


def _optional_workbook_text(value: object) -> str | None:
    if value is None:
        return None
    text = " ".join(str(value).split())
    return text or None


def _parse_current_inventory_workbook(
    workbook_path: Path,
    *,
    expected_sha256: str,
) -> list[dict[str, Any]]:
    observed_sha256 = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    if observed_sha256 != expected_sha256:
        raise ValueError("inventory workbook hash does not match the V5 authority")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if _CURRENT_MASTER_SHEET not in workbook.sheetnames:
            raise ValueError("inventory workbook is missing Current Inventory Master")
        sheet = workbook[_CURRENT_MASTER_SHEET]
        header_values = next(
            sheet.iter_rows(min_row=5, max_row=5, values_only=True),
            None,
        )
        if header_values is None:
            raise ValueError("inventory workbook is missing the row-5 header")
        header_indexes = {
            str(value).strip(): index
            for index, value in enumerate(header_values)
            if value is not None and str(value).strip()
        }
        missing_headers = tuple(
            header for header in _CURRENT_MASTER_HEADERS if header not in header_indexes
        )
        if missing_headers:
            raise ValueError(
                "inventory workbook is missing required headers: "
                + ", ".join(missing_headers)
            )

        records: list[dict[str, Any]] = []
        for source_row, values in enumerate(
            sheet.iter_rows(min_row=6, values_only=True),
            start=6,
        ):
            if not any(value is not None and str(value).strip() for value in values):
                continue
            record = {
                field_name: _optional_workbook_text(
                    values[header_indexes[header]]
                    if header_indexes[header] < len(values)
                    else None
                )
                for header, field_name in _CURRENT_MASTER_HEADERS.items()
            }
            record["source_row"] = source_row
            records.append(record)
        return records
    finally:
        workbook.close()


def _load_execution_inventory_catalog(
    catalog_path: str | None,
    workbook_path: str,
) -> ComplexityInventoryCatalog:
    path = (
        Path(catalog_path)
        if catalog_path is not None
        else Path(_complexity_inventory._CATALOG_PATH)
    )
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, Mapping):
        raise TypeError("complexity inventory catalog root must be a mapping")
    authority = value.get("authority")
    if not isinstance(authority, Mapping):
        raise TypeError("authority must be a mapping")
    expected_sha256 = _text(authority.get("workbook_sha256"), "workbook_sha256")
    current_records = _parse_current_inventory_workbook(
        Path(workbook_path),
        expected_sha256=expected_sha256,
    )
    runtime_value = dict(value)
    runtime_value["current_records"] = current_records
    return ComplexityInventoryCatalog(runtime_value)


def _text_tuple(values: tuple[str, ...], field_name: str) -> tuple[str, ...]:
    normalized = tuple(_text(value, field_name) for value in values)
    if len(normalized) != len(set(normalized)):
        raise ValueError(f"{field_name} must not contain duplicates")
    return normalized


@dataclass(frozen=True, slots=True)
class ComparisonClosureV1:
    """Complete one-comparison decision contract for an architectural delta."""

    rejected_alternative: str
    compliant_treatment: str
    primary_endpoints: tuple[str, ...]
    failure_endpoints: tuple[str, ...]
    changed_factor: str
    constant_constraints: tuple[str, ...]
    blinding_rule: str
    order_rule: str
    time_windows: tuple[str, ...]
    accept_rule: str
    reject_rule: str

    def __post_init__(self) -> None:
        for name in (
            "rejected_alternative",
            "compliant_treatment",
            "changed_factor",
            "blinding_rule",
            "order_rule",
            "accept_rule",
            "reject_rule",
        ):
            object.__setattr__(self, name, _text(getattr(self, name), name))
        for name in (
            "primary_endpoints",
            "failure_endpoints",
            "constant_constraints",
            "time_windows",
        ):
            values = _text_tuple(tuple(getattr(self, name)), name)
            if not values:
                raise ValueError(f"{name} must be nonempty")
            object.__setattr__(self, name, values)
        if self.rejected_alternative == self.compliant_treatment:
            raise ValueError("comparison alternatives must be distinct")

    def as_dict(self) -> dict[str, object]:
        return {
            "schema_version": "comparison_closure_v1",
            "rejected_alternative": self.rejected_alternative,
            "compliant_treatment": self.compliant_treatment,
            "primary_endpoints": list(self.primary_endpoints),
            "failure_endpoints": list(self.failure_endpoints),
            "changed_factor": self.changed_factor,
            "constant_constraints": list(self.constant_constraints),
            "blinding_rule": self.blinding_rule,
            "order_rule": self.order_rule,
            "time_windows": list(self.time_windows),
            "accept_rule": self.accept_rule,
            "reject_rule": self.reject_rule,
        }

    @property
    def closure_sha256(self) -> str:
        return sha256_hex(canonical_json_bytes(self.as_dict()))


@dataclass(frozen=True, slots=True)
class ArchitecturalDeltaCandidate:
    candidate_id: str
    material: str
    family: ArchitecturalDeltaFamily
    kind: ArchitecturalDeltaKind
    priority_rank: int
    target_role: str
    nonredundancy_evidence: str
    loss_if_omitted: str
    failure_mode: str
    controlled_arms: tuple[str, ...]
    evidence_refs: tuple[str, ...]
    count_based: bool = False
    redundant_with_current_build: bool = False
    primary_role: bool = False
    target_explicitly_names_material: bool = False
    distinct_role_refs: tuple[str, ...] = ()
    pairwise_nonredundancy_refs: tuple[str, ...] = ()
    exception_justification: str | None = None
    causal_design_sha256: str | None = None
    nary_candidate: NaryInteractionCandidate | None = None
    uniqueness_evidence_refs: tuple[str, ...] = ()
    comparison_closure: ComparisonClosureV1 | None = None

    def __post_init__(self) -> None:
        for name in (
            "candidate_id",
            "material",
            "target_role",
            "nonredundancy_evidence",
            "loss_if_omitted",
            "failure_mode",
        ):
            object.__setattr__(self, name, _text(getattr(self, name), name))
        object.__setattr__(self, "family", ArchitecturalDeltaFamily(self.family))
        object.__setattr__(self, "kind", ArchitecturalDeltaKind(self.kind))
        if isinstance(self.priority_rank, bool) or self.priority_rank < 1:
            raise ValueError("priority_rank must be a positive integer")
        for name in (
            "count_based",
            "redundant_with_current_build",
            "primary_role",
            "target_explicitly_names_material",
        ):
            if not isinstance(getattr(self, name), bool):
                raise TypeError(f"{name} must be boolean")
        if len(self.controlled_arms) < 2:
            raise ValueError("controlled_arms requires at least two arms")
        object.__setattr__(
            self,
            "controlled_arms",
            _text_tuple(tuple(self.controlled_arms), "controlled_arms"),
        )
        object.__setattr__(
            self,
            "evidence_refs",
            _text_tuple(tuple(self.evidence_refs), "evidence_refs"),
        )
        object.__setattr__(
            self,
            "distinct_role_refs",
            _text_tuple(tuple(self.distinct_role_refs), "distinct_role_refs"),
        )
        object.__setattr__(
            self,
            "pairwise_nonredundancy_refs",
            _text_tuple(
                tuple(self.pairwise_nonredundancy_refs),
                "pairwise_nonredundancy_refs",
            ),
        )
        object.__setattr__(
            self,
            "uniqueness_evidence_refs",
            _text_tuple(
                tuple(self.uniqueness_evidence_refs),
                "uniqueness_evidence_refs",
            ),
        )
        if self.exception_justification is not None:
            object.__setattr__(
                self,
                "exception_justification",
                _text(self.exception_justification, "exception_justification"),
            )
        if self.causal_design_sha256 is not None:
            object.__setattr__(
                self,
                "causal_design_sha256",
                _sha256(self.causal_design_sha256, "causal_design_sha256"),
            )
        if self.nary_candidate is not None and not isinstance(
            self.nary_candidate, NaryInteractionCandidate
        ):
            raise TypeError("nary_candidate must be a NaryInteractionCandidate")
        if self.comparison_closure is not None and not isinstance(
            self.comparison_closure, ComparisonClosureV1
        ):
            raise TypeError("comparison_closure must be a ComparisonClosureV1")


@dataclass(frozen=True, slots=True)
class ArchitecturalDeltaRequest:
    target_identity: str
    ideal_formula_ref: str
    current_build_ref: str
    formula_lineage_sha256: str
    candidates: tuple[ArchitecturalDeltaCandidate, ...]
    no_change_reason: str

    def __post_init__(self) -> None:
        for name in (
            "target_identity",
            "ideal_formula_ref",
            "current_build_ref",
            "no_change_reason",
        ):
            object.__setattr__(self, name, _text(getattr(self, name), name))
        object.__setattr__(
            self,
            "formula_lineage_sha256",
            _sha256(self.formula_lineage_sha256, "formula_lineage_sha256"),
        )
        object.__setattr__(self, "candidates", tuple(self.candidates))
        if any(
            not isinstance(candidate, ArchitecturalDeltaCandidate)
            for candidate in self.candidates
        ):
            raise TypeError(
                "candidates must contain ArchitecturalDeltaCandidate values"
            )
        candidate_ids = tuple(candidate.candidate_id for candidate in self.candidates)
        if len(candidate_ids) != len(set(candidate_ids)):
            raise ValueError("candidate IDs must be unique")


@dataclass(frozen=True, slots=True)
class ArchitecturalDeltaResult:
    state: ArchitecturalDeltaState
    target_identity: str
    ideal_formula_ref: str
    current_build_ref: str
    formula_lineage_sha256: str
    inventory_workbook_sha256: str
    inventory_source_row_count: int
    selected_candidate: ArchitecturalDeltaCandidate | None
    inventory_projection: InventoryProjection | None
    controlled_arms: tuple[str, ...]
    blockers: tuple[str, ...]
    next_comparison: str | None
    empirical_authority: bool = field(default=False, init=False)
    formula_mutation_authorized: bool = field(default=False, init=False)
    physical_execution_authorized: bool = field(default=False, init=False)
    purchase_authority: bool = field(default=False, init=False)
    sensory_authority: bool = field(default=False, init=False)
    safety_authority: bool = field(default=False, init=False)
    release_authority: bool = field(default=False, init=False)


@dataclass(frozen=True, slots=True)
class ArchitecturalEvidenceDeltaResultV2:
    """V2 augmentation disposition layered over the compatible V1 result."""

    architectural_result: ArchitecturalDeltaResult
    comparison_closure: ComparisonClosureV1 | None
    receipt: EvidenceDeltaReceiptV1

    def __post_init__(self) -> None:
        if not isinstance(self.architectural_result, ArchitecturalDeltaResult):
            raise TypeError("architectural_result must be an ArchitecturalDeltaResult")
        if self.comparison_closure is not None and not isinstance(
            self.comparison_closure, ComparisonClosureV1
        ):
            raise TypeError("comparison_closure must be a ComparisonClosureV1")
        if not isinstance(self.receipt, EvidenceDeltaReceiptV1):
            raise TypeError("receipt must be an EvidenceDeltaReceiptV1")


def evaluate_architectural_delta(
    request: ArchitecturalDeltaRequest,
    *,
    inventory_catalog_path: str | None = None,
    inventory_workbook_path: str | None = None,
) -> ArchitecturalDeltaResult:
    """Evaluate one target-defined delta against the supplied V5 authority."""

    if not isinstance(request, ArchitecturalDeltaRequest):
        raise TypeError("request must be an ArchitecturalDeltaRequest")
    if inventory_workbook_path is None:
        inventory_workbook_path = os.environ.get(
            "PERFUME_COMPLEXITY_INVENTORY_WORKBOOK"
        )
    if not inventory_workbook_path or not inventory_workbook_path.strip():
        raise ValueError(
            "authoritative inventory workbook path is required for every "
            "Architectural Delta execution"
        )
    catalog = _load_execution_inventory_catalog(
        inventory_catalog_path,
        inventory_workbook_path,
    )
    common = {
        "target_identity": request.target_identity,
        "ideal_formula_ref": request.ideal_formula_ref,
        "current_build_ref": request.current_build_ref,
        "formula_lineage_sha256": request.formula_lineage_sha256,
        "inventory_workbook_sha256": catalog.workbook_sha256,
        "inventory_source_row_count": len(catalog.current_records),
    }
    if not request.candidates:
        return ArchitecturalDeltaResult(
            state=ArchitecturalDeltaState.NO_CHANGE,
            selected_candidate=None,
            inventory_projection=None,
            controlled_arms=(),
            blockers=(),
            next_comparison=None,
            **common,
        )

    eligible: list[ArchitecturalDeltaCandidate] = []
    blockers: list[str] = []
    for candidate in request.candidates:
        candidate_blockers = _candidate_blockers(candidate)
        if candidate_blockers:
            blockers.extend(candidate_blockers)
        else:
            eligible.append(candidate)
    if not eligible:
        return ArchitecturalDeltaResult(
            state=ArchitecturalDeltaState.HOLD,
            selected_candidate=None,
            inventory_projection=None,
            controlled_arms=(),
            blockers=tuple(blockers),
            next_comparison=None,
            **common,
        )

    eligible.sort(key=lambda candidate: (candidate.priority_rank, candidate.candidate_id))
    first_rank = eligible[0].priority_rank
    top_ranked = tuple(
        candidate for candidate in eligible if candidate.priority_rank == first_rank
    )
    if len(top_ranked) != 1:
        return ArchitecturalDeltaResult(
            state=ArchitecturalDeltaState.HOLD,
            selected_candidate=None,
            inventory_projection=None,
            controlled_arms=(),
            blockers=("target architecture gives multiple candidates the same priority",),
            next_comparison=None,
            **common,
        )

    selected = top_ranked[0]
    projection = catalog.project(selected.material)
    return ArchitecturalDeltaResult(
        state=ArchitecturalDeltaState.PROPOSED,
        selected_candidate=selected,
        inventory_projection=projection,
        controlled_arms=selected.controlled_arms,
        blockers=(),
        next_comparison=(
            f"Compare {selected.controlled_arms[0]} with "
            f"{selected.controlled_arms[1]} for {selected.target_role}."
        ),
        **common,
    )


_ARCHITECTURAL_POLICY = {
    "schema_version": "architectural_evidence_delta_policy_v2",
    "target_first": True,
    "ingredient_count_reward": False,
    "zero_or_one_intervention": True,
    "neroli_default": "SUPPORT_ONLY",
    "musk_default": "ZERO_OR_ONE",
    "exception_only_musks": ["Tonalide", "Macrolide", "Musk Ketone"],
    "nary_requires_isolated_arms": True,
    "authority": "EVIDENCE_DESIGN_ONLY",
}
_ARCHITECTURAL_POLICY_SHA256 = sha256_hex(
    canonical_json_bytes(_ARCHITECTURAL_POLICY)
)


def _request_evidence_payload(request: ArchitecturalDeltaRequest) -> dict[str, object]:
    return {
        "target_identity": request.target_identity,
        "ideal_formula_ref": request.ideal_formula_ref,
        "current_build_ref": request.current_build_ref,
        "formula_lineage_sha256": request.formula_lineage_sha256,
        "no_change_reason": request.no_change_reason,
        "candidates": [
            {
                "candidate_id": candidate.candidate_id,
                "material": candidate.material,
                "family": candidate.family.value,
                "kind": candidate.kind.value,
                "priority_rank": candidate.priority_rank,
                "target_role": candidate.target_role,
                "nonredundancy_evidence": candidate.nonredundancy_evidence,
                "controlled_arms": list(candidate.controlled_arms),
                "evidence_refs": list(candidate.evidence_refs),
                "uniqueness_evidence_refs": list(
                    candidate.uniqueness_evidence_refs
                ),
                "comparison_closure": (
                    None
                    if candidate.comparison_closure is None
                    else candidate.comparison_closure.as_dict()
                ),
            }
            for candidate in request.candidates
        ],
    }


def _source_binding_hashes(
    candidate: ArchitecturalDeltaCandidate,
) -> tuple[str, ...]:
    refs = (*candidate.evidence_refs, *candidate.uniqueness_evidence_refs)
    return tuple(
        sorted(
            {
                value.casefold()
                for value in refs
                if len(value) == 64
                and all(character in "0123456789abcdefABCDEF" for character in value)
            }
        )
    )


def _v2_reason_codes(blockers: tuple[str, ...]) -> tuple[str, ...]:
    codes: list[str] = []
    joined = " ".join(blockers).casefold()
    if "comparison closure is required" in joined:
        codes.append("COMPARISON_CLOSURE_MISSING")
    elif "comparison closure" in joined:
        codes.append("COMPARISON_CLOSURE_INVALID")
    if "unique" in joined and "evidence" in joined:
        codes.append("UNIQUENESS_EVIDENCE_MISSING")
    if "source binding" in joined:
        codes.append("SOURCE_BINDING_MISSING")
    if not codes:
        codes.append("ARCHITECTURAL_POLICY_HOLD")
    return tuple(codes)


def evaluate_architectural_evidence_delta(
    request: ArchitecturalDeltaRequest,
    *,
    inventory_catalog_path: str | None = None,
    inventory_workbook_path: str | None = None,
) -> ArchitecturalEvidenceDeltaResultV2:
    """Emit one closed decision delta, explicit abstention, or a precise hold."""

    result = evaluate_architectural_delta(
        request,
        inventory_catalog_path=inventory_catalog_path,
        inventory_workbook_path=inventory_workbook_path,
    )
    request_payload = _request_evidence_payload(request)
    input_sha256 = sha256_hex(canonical_json_bytes(request_payload))
    evidence_sha256 = sha256_hex(
        canonical_json_bytes(
            {
                "candidate_evidence": [
                    {
                        "candidate_id": candidate.candidate_id,
                        "evidence_refs": list(candidate.evidence_refs),
                        "uniqueness_evidence_refs": list(
                            candidate.uniqueness_evidence_refs
                        ),
                    }
                    for candidate in request.candidates
                ]
            }
        )
    )
    exact_scope = (
        f"{request.target_identity}/{request.ideal_formula_ref}/"
        f"{request.current_build_ref}"
    )
    if result.state is ArchitecturalDeltaState.NO_CHANGE:
        receipt = no_augmentation_receipt(
            module_id="architectural_delta",
            exact_scope=exact_scope,
            input_sha256=input_sha256,
            evidence_sha256=evidence_sha256,
            policy_sha256=_ARCHITECTURAL_POLICY_SHA256,
            reasons=("NO_TARGET_DEFICIENCY",),
        )
        return ArchitecturalEvidenceDeltaResultV2(result, None, receipt)

    selected = result.selected_candidate
    blockers = list(result.blockers)
    closure = None if selected is None else selected.comparison_closure
    if result.state is ArchitecturalDeltaState.PROPOSED and selected is not None:
        if closure is None:
            blockers.append(
                f"{selected.candidate_id}: comparison closure is required"
            )
        else:
            arms = set(selected.controlled_arms)
            if {
                closure.rejected_alternative,
                closure.compliant_treatment,
            }.difference(arms):
                blockers.append(
                    f"{selected.candidate_id}: comparison closure alternatives "
                    "must be controlled arms"
                )
            if not any(
                "TOTAL" in constraint.upper()
                for constraint in closure.constant_constraints
            ):
                blockers.append(
                    f"{selected.candidate_id}: comparison closure requires a "
                    "constant-total constraint"
                )
        source_bindings = _source_binding_hashes(selected)
        if not source_bindings:
            blockers.append(f"{selected.candidate_id}: source binding is required")
    else:
        source_bindings = ()

    if blockers or selected is None or closure is None:
        blocker_tuple = tuple(blockers) or ("No eligible intervention remains.",)
        receipt = hold_receipt(
            module_id="architectural_delta",
            exact_scope=exact_scope,
            input_sha256=input_sha256,
            evidence_sha256=evidence_sha256,
            policy_sha256=_ARCHITECTURAL_POLICY_SHA256,
            source_binding_sha256=source_bindings,
            reasons=_v2_reason_codes(blocker_tuple),
            blockers=blocker_tuple,
            next_action=(
                "CLOSE_COMPARISON"
                if any("comparison closure" in value.casefold() for value in blockers)
                else None
            ),
        )
        return ArchitecturalEvidenceDeltaResultV2(result, closure, receipt)

    projection = result.inventory_projection
    observed_facts = (
        f"target_identity={request.target_identity}",
        f"ideal_formula_ref={request.ideal_formula_ref}",
        f"current_build_ref={request.current_build_ref}",
        f"inventory_status={projection.availability.value if projection else 'UNKNOWN'}",
    )
    delta = DecisionDeltaV1(
        delta_id=f"ARCHITECTURAL:{selected.candidate_id}",
        decision_effect=(
            f"Run one closed comparison of {closure.rejected_alternative} against "
            f"{closure.compliant_treatment}."
        ),
        observed_facts=observed_facts,
        derived_calculations=(
            "changed_factors=1",
            f"controlled_arms={len(selected.controlled_arms)}",
            f"inventory_source_rows={result.inventory_source_row_count}",
            f"closure_sha256={closure.closure_sha256}",
        ),
        hypotheses=(
            selected.nonredundancy_evidence,
            f"omission_loss={selected.loss_if_omitted}",
        ),
        forbidden_inferences=(
            "No composition-derived liking, sensory truth, safety, purchase, or release claim is permitted.",
            "The proposed arm is an experiment, not a formula mutation authorization.",
        ),
    )
    receipt = EvidenceDeltaReceiptV1(
        module_id="architectural_delta",
        exact_scope=exact_scope,
        state=EvidenceAugmentationState.AUGMENT,
        input_sha256=input_sha256,
        evidence_sha256=evidence_sha256,
        policy_sha256=_ARCHITECTURAL_POLICY_SHA256,
        source_binding_sha256=source_bindings,
        reason_codes=("ONE_CLOSED_COMPARISON",),
        delta=delta,
        blockers=(),
        next_action=(
            f"COMPARE:{closure.rejected_alternative}:"
            f"{closure.compliant_treatment}"
        ),
    )
    return ArchitecturalEvidenceDeltaResultV2(result, closure, receipt)


_EXCEPTION_ONLY_MUSKS = ("tonalide", "macrolide", "musk ketone")


def _candidate_blockers(
    candidate: ArchitecturalDeltaCandidate,
) -> tuple[str, ...]:
    blockers: list[str] = []
    if candidate.count_based:
        blockers.append(
            f"{candidate.candidate_id}: ingredient count is not complexity evidence"
        )
    if candidate.redundant_with_current_build:
        blockers.append(
            f"{candidate.candidate_id}: addition is redundant with the current build"
        )
    uniqueness_claim = (
        "unique" in candidate.target_role.casefold()
        or "unique" in candidate.nonredundancy_evidence.casefold()
    )
    if uniqueness_claim and not candidate.uniqueness_evidence_refs:
        blockers.append(
            f"{candidate.candidate_id}: unique role requires uniqueness evidence"
        )
    material = candidate.material.casefold()
    if (
        candidate.family is ArchitecturalDeltaFamily.CITRUS
        and "neroli" in material
        and candidate.primary_role
        and not candidate.target_explicitly_names_material
    ):
        blockers.append(
            f"{candidate.candidate_id}: Neroli is support-only unless the target "
            "explicitly requires it as primary"
        )
    if (
        candidate.family is ArchitecturalDeltaFamily.MUSK
        and any(name in material for name in _EXCEPTION_ONLY_MUSKS)
        and candidate.exception_justification is None
    ):
        blockers.append(
            f"{candidate.candidate_id}: Tonalide, Macrolide, and Musk Ketone are "
            "exception-only"
        )
    if candidate.kind is ArchitecturalDeltaKind.NARY_DESIGN:
        role_count = len(candidate.distinct_role_refs)
        required_pairs = role_count * (role_count - 1) // 2
        if role_count < 2:
            blockers.append(
                f"{candidate.candidate_id}: n-ary design requires distinct roles"
            )
        if len(candidate.pairwise_nonredundancy_refs) < required_pairs:
            blockers.append(
                f"{candidate.candidate_id}: n-ary design lacks complete pairwise "
                "nonredundancy isolates"
            )
        if role_count == 2:
            factors = tuple(
                part.strip().upper().replace(" ", "_")
                for part in candidate.material.split("+")
                if part.strip()
            )
            expected_arms = (
                "CONTROL",
                *factors,
                "_X_".join(factors),
            )
            if len(factors) != 2 or candidate.controlled_arms != expected_arms:
                blockers.append(
                    f"{candidate.candidate_id}: two-factor interaction requires "
                    "exactly four factorial arms"
                )
            if candidate.causal_design_sha256 is None:
                blockers.append(
                    f"{candidate.candidate_id}: two-factor interaction requires a "
                    "hash-bound causal design"
                )
        elif candidate.nary_candidate is None:
            blockers.append(
                f"{candidate.candidate_id}: existing n-ary interaction contract is required"
            )
        else:
            assessment = evaluate_nary_interaction(candidate.nary_candidate)
            if assessment.state is not NaryAssessmentState.DESIGN_READY:
                blockers.append(
                    f"{candidate.candidate_id}: n-ary interaction contract is "
                    f"{assessment.state.value}"
                )
                blockers.extend(
                    f"{candidate.candidate_id}: {blocker}"
                    for blocker in assessment.blockers
                )
    return tuple(blockers)


__all__ = [
    "ArchitecturalEvidenceDeltaResultV2",
    "ArchitecturalDeltaRequest",
    "ArchitecturalDeltaResult",
    "ArchitecturalDeltaCandidate",
    "ArchitecturalDeltaFamily",
    "ArchitecturalDeltaKind",
    "ArchitecturalDeltaState",
    "ComparisonClosureV1",
    "evaluate_architectural_delta",
    "evaluate_architectural_evidence_delta",
]
