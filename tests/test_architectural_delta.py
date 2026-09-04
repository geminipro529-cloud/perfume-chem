from __future__ import annotations

import hashlib
import json
import os
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook

from engine.perception.architectural_delta import (
    ArchitecturalDeltaCandidate,
    ArchitecturalDeltaFamily,
    ArchitecturalDeltaKind,
    ArchitecturalDeltaRequest,
    ArchitecturalDeltaState,
    evaluate_architectural_delta,
)
from engine.perception.complexity_inventory import (
    InventoryAvailability,
    StockReadiness,
)
from engine.scientific_validation.complexity_design_contracts import (
    NaryEvidenceState,
    NaryInteractionCandidate,
    NaryParticipant,
)
from engine.scientific_validation.complexity_model_admission import OAVGateBinding


@pytest.fixture(autouse=True)
def _runtime_inventory_authority(tmp_path, monkeypatch) -> None:
    source_catalog_path = (
        Path(__file__).parents[1]
        / "data"
        / "governance"
        / "complexity_inventory_catalog_v1.json"
    )
    catalog = json.loads(source_catalog_path.read_text(encoding="utf-8"))
    workbook_path = tmp_path / "runtime-inventory-v5.xlsx"
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Current Inventory Master")
    sheet.append(["title"])
    sheet.append(["authority note"])
    sheet.append([])
    sheet.append([])
    sheet.append(
        [
            "Baseline #",
            "Canonical material",
            "Status",
            "Actual stock(s)",
            "Can prepare",
            "Family",
            "Prior listed stock",
            "Prior active fraction",
            "Formula count",
            "Row uses",
            "Affected target IDs",
            "Alias / non-equivalent",
            "Formula-use policy",
            "User note",
        ]
    )
    for record in catalog["current_records"]:
        sheet.append(
            [
                None,
                record["canonical_material"],
                record["status"],
                record["actual_stocks"],
                record["can_prepare"],
                record["family"],
                None,
                None,
                0,
                0,
                None,
                record["alias_or_non_equivalent"],
                record["formula_use_policy"],
                record["user_note"],
            ]
        )
    workbook.save(workbook_path)
    workbook.close()
    workbook_sha256 = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    catalog["authority"]["workbook_sha256"] = workbook_sha256
    catalog_path = tmp_path / "runtime-catalog.json"
    catalog_path.write_text(json.dumps(catalog), encoding="utf-8")

    from engine.perception import complexity_inventory

    monkeypatch.setattr(
        complexity_inventory,
        "_EXPECTED_WORKBOOK_SHA256",
        workbook_sha256,
    )
    monkeypatch.setattr(complexity_inventory, "_CATALOG_PATH", catalog_path)
    monkeypatch.setenv(
        "PERFUME_COMPLEXITY_INVENTORY_WORKBOOK",
        str(workbook_path),
    )


def test_empty_target_defined_candidate_set_returns_authority_free_no_change() -> None:
    request = ArchitecturalDeltaRequest(
        target_identity="Austere iris with deliberate negative space",
        ideal_formula_ref="formula:iris:ideal:v1",
        current_build_ref="formula:iris:inventory:v1",
        formula_lineage_sha256="1" * 64,
        candidates=(),
        no_change_reason="The target is already complete through precise simplicity.",
    )

    result = evaluate_architectural_delta(request)

    assert result.state is ArchitecturalDeltaState.NO_CHANGE
    assert result.ideal_formula_ref == "formula:iris:ideal:v1"
    assert result.current_build_ref == "formula:iris:inventory:v1"
    expected_workbook_sha256 = hashlib.sha256(
        Path(os.environ["PERFUME_COMPLEXITY_INVENTORY_WORKBOOK"]).read_bytes()
    ).hexdigest()
    assert result.inventory_workbook_sha256 == expected_workbook_sha256
    assert result.selected_candidate is None
    assert result.next_comparison is None
    assert result.formula_mutation_authorized is False
    assert result.physical_execution_authorized is False
    assert result.purchase_authority is False
    assert result.sensory_authority is False
    assert result.safety_authority is False
    assert result.release_authority is False


def test_execution_requires_authoritative_workbook_path(monkeypatch) -> None:
    monkeypatch.delenv("PERFUME_COMPLEXITY_INVENTORY_WORKBOOK")

    with pytest.raises(ValueError, match="inventory workbook path"):
        evaluate_architectural_delta(
            ArchitecturalDeltaRequest(
                target_identity="Precise target",
                ideal_formula_ref="formula:target:ideal:v1",
                current_build_ref="formula:target:inventory:v1",
                formula_lineage_sha256="a" * 64,
                candidates=(),
                no_change_reason="No justified delta remains.",
            )
        )


def _candidate(
    *,
    candidate_id: str,
    material: str,
    family: ArchitecturalDeltaFamily = ArchitecturalDeltaFamily.GENERAL,
    kind: ArchitecturalDeltaKind = ArchitecturalDeltaKind.ADDITION,
    priority_rank: int = 1,
    target_role: str = "increase target-defined depth without changing identity",
    **overrides: object,
) -> ArchitecturalDeltaCandidate:
    values = {
        "candidate_id": candidate_id,
        "material": material,
        "family": family,
        "kind": kind,
        "priority_rank": priority_rank,
        "target_role": target_role,
        "nonredundancy_evidence": "The candidate supplies a missing target function.",
        "loss_if_omitted": "The target loses its intended depth transition.",
        "failure_mode": "An overdose would blur the target identity.",
        "controlled_arms": ("baseline", f"single-{candidate_id}"),
        "evidence_refs": (f"target:{candidate_id}",),
    }
    values.update(overrides)
    return ArchitecturalDeltaCandidate(**values)


def _request(*candidates: ArchitecturalDeltaCandidate) -> ArchitecturalDeltaRequest:
    return ArchitecturalDeltaRequest(
        target_identity="Target-defined depth without ingredient-count reward",
        ideal_formula_ref="formula:target:ideal:v1",
        current_build_ref="formula:target:inventory:v1",
        formula_lineage_sha256="2" * 64,
        candidates=candidates,
        no_change_reason="No justified architectural delta remains.",
    )


def _nary_contract() -> NaryInteractionCandidate:
    formula_sha256 = "3" * 64
    return NaryInteractionCandidate(
        interaction_id="nary-musk-depth",
        evidence_state=NaryEvidenceState.EXPERIMENT_DESIGN,
        participants=(
            NaryParticipant("Habanolide", "radiance", Decimal("0.4")),
            NaryParticipant("Romandolide", "texture", Decimal("0.35")),
            NaryParticipant("Ambrettolide 10%", "skin continuity", Decimal("0.25")),
        ),
        formula_sha256=formula_sha256,
        formula_signature_sha256="4" * 64,
        matrix_sha256="5" * 64,
        pairwise_evidence_refs=("hab-rom", "hab-amb", "rom-amb"),
        causal_isolate_sha256="6" * 64,
        experiment_design_sha256="7" * 64,
        oav_binding=OAVGateBinding(
            formula_sha256=formula_sha256,
            dose_receipt_sha256="8" * 64,
            oav_result_sha256="9" * 64,
            quantitative_ppm_status="PASS",
            odt_authority_status="PASS",
            odt_coverage_status="PASS",
            natural_composite_coverage_status="PASS",
            headspace_scope_status="PASS",
            receipt_binding_status="BOUND_GATE_RECEIPT",
            strict_oav_status="NOT_TESTED",
            pre_mix_gate_status="PASS",
            planned_active_equivalence_status="PASS",
            formula_is_revision=False,
        ),
    )


def test_selects_one_target_ranked_candidate_and_binds_current_inventory() -> None:
    result = evaluate_architectural_delta(
        _request(
            _candidate(
                candidate_id="soft-volume",
                material="Romandolide",
                family=ArchitecturalDeltaFamily.MUSK,
                priority_rank=2,
            ),
            _candidate(
                candidate_id="radiant-dryness",
                material="Habanolide",
                family=ArchitecturalDeltaFamily.MUSK,
                priority_rank=1,
            ),
        )
    )

    assert result.state is ArchitecturalDeltaState.PROPOSED
    assert result.selected_candidate is not None
    assert result.selected_candidate.candidate_id == "radiant-dryness"
    assert result.inventory_projection is not None
    assert result.inventory_projection.availability is InventoryAvailability.OWNED
    assert result.inventory_projection.stock_readiness is StockReadiness.EXACT_STOCK_IDENTIFIED
    assert result.controlled_arms == ("baseline", "single-radiant-dryness")
    assert "baseline" in (result.next_comparison or "")


@pytest.mark.parametrize(
    ("material", "availability", "readiness"),
    (
        ("Habanolide", InventoryAvailability.OWNED, StockReadiness.EXACT_STOCK_IDENTIFIED),
        ("Clearwood", InventoryAvailability.OUT_OF_STOCK, StockReadiness.NOT_BUILDABLE),
        (
            "Ambrettolide 10%",
            InventoryAvailability.PLANNED_ACQUISITION,
            StockReadiness.PROCUREMENT_PENDING,
        ),
        (
            "Ethylene Brassylate",
            InventoryAvailability.MISSING,
            StockReadiness.NOT_BUILDABLE,
        ),
        (
            "Apple Accord 10%",
            InventoryAvailability.PREPARABLE_NOT_MIXED,
            StockReadiness.PREPARATION_REQUIRED,
        ),
        (
            "Basil EO ct. Methyl Chavicol 10%",
            InventoryAvailability.VERIFY_FIRST,
            StockReadiness.STOCK_DETAIL_OPEN,
        ),
        (
            "DO NOT USE — Iris FTEC",
            InventoryAvailability.FORBIDDEN,
            StockReadiness.NOT_BUILDABLE,
        ),
        (
            "Imaginary Material X",
            InventoryAvailability.UNLISTED,
            StockReadiness.NOT_BUILDABLE,
        ),
    ),
)
def test_every_inventory_status_is_reported_without_redefining_the_target(
    material: str,
    availability: InventoryAvailability,
    readiness: StockReadiness,
) -> None:
    result = evaluate_architectural_delta(
        _request(_candidate(candidate_id="status-check", material=material))
    )

    assert result.state is ArchitecturalDeltaState.PROPOSED
    assert result.ideal_formula_ref == "formula:target:ideal:v1"
    assert result.current_build_ref == "formula:target:inventory:v1"
    assert result.inventory_projection is not None
    assert result.inventory_projection.availability is availability
    assert result.inventory_projection.stock_readiness is readiness


def test_execution_reparses_authoritative_workbook_instead_of_stale_catalog(
    tmp_path, monkeypatch
) -> None:
    workbook_path = tmp_path / "inventory-v5.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Current Inventory Master"
    sheet.append(["title"])
    sheet.append(["authority note"])
    sheet.append([])
    sheet.append([])
    sheet.append(
        [
            "Baseline #",
            "Canonical material",
            "Status",
            "Actual stock(s)",
            "Can prepare",
            "Family",
            "Prior listed stock",
            "Prior active fraction",
            "Formula count",
            "Row uses",
            "Affected target IDs",
            "Alias / non-equivalent",
            "Formula-use policy",
            "User note",
        ]
    )
    sheet.append(
        [
            1,
            "Runtime Workbook Material",
            "HAVE — NEAT",
            "Runtime Workbook Material neat/as supplied",
            None,
            "test family",
            None,
            None,
            0,
            0,
            None,
            None,
            "Use only with exact stock accounting.",
            "Fixture authority row.",
        ]
    )
    workbook.save(workbook_path)
    workbook.close()
    workbook_sha256 = hashlib.sha256(workbook_path.read_bytes()).hexdigest()

    catalog_path = tmp_path / "catalog.json"
    catalog = json.loads(
        (
            Path(__file__).parents[1]
            / "data"
            / "governance"
            / "complexity_inventory_catalog_v1.json"
        ).read_text(encoding="utf-8")
    )
    catalog["authority"]["workbook_sha256"] = workbook_sha256
    catalog["authority"]["current_record_count"] = 1
    catalog["current_records"] = [
        {
            "source_row": 6,
            "canonical_material": "Stale Catalog Material",
            "status": "OUT OF STOCK",
            "actual_stocks": None,
            "can_prepare": None,
            "family": "stale family",
            "alias_or_non_equivalent": None,
            "formula_use_policy": "Stale projection must not win.",
            "user_note": None,
        }
    ]
    catalog_path.write_text(json.dumps(catalog), encoding="utf-8")

    from engine.perception import complexity_inventory

    monkeypatch.setattr(
        complexity_inventory,
        "_EXPECTED_WORKBOOK_SHA256",
        workbook_sha256,
    )
    result = evaluate_architectural_delta(
        _request(
            _candidate(
                candidate_id="runtime-workbook",
                material="Runtime Workbook Material",
            )
        ),
        inventory_catalog_path=str(catalog_path),
        inventory_workbook_path=str(workbook_path),
    )

    assert result.state is ArchitecturalDeltaState.PROPOSED
    assert result.inventory_projection is not None
    assert result.inventory_projection.canonical_material == "Runtime Workbook Material"
    assert result.inventory_projection.availability is InventoryAvailability.OWNED
    assert result.inventory_projection.source_row == 6


def test_count_rationale_and_generic_neroli_primary_are_held() -> None:
    count_result = evaluate_architectural_delta(
        _request(
            _candidate(
                candidate_id="more-rows",
                material="Habanolide",
                count_based=True,
            )
        )
    )
    neroli_result = evaluate_architectural_delta(
        _request(
            _candidate(
                candidate_id="generic-primary-neroli",
                material="Neroli EO 10%",
                family=ArchitecturalDeltaFamily.CITRUS,
                primary_role=True,
                target_explicitly_names_material=False,
            )
        )
    )

    assert count_result.state is ArchitecturalDeltaState.HOLD
    assert any("ingredient count" in blocker for blocker in count_result.blockers)
    assert neroli_result.state is ArchitecturalDeltaState.HOLD
    assert any("support-only" in blocker for blocker in neroli_result.blockers)


def test_redundant_addition_is_held_even_when_it_is_owned() -> None:
    result = evaluate_architectural_delta(
        _request(
            _candidate(
                candidate_id="duplicate-volume",
                material="Romandolide",
                redundant_with_current_build=True,
            )
        )
    )

    assert result.state is ArchitecturalDeltaState.HOLD
    assert any("redundant" in blocker for blocker in result.blockers)


def test_exception_musk_requires_an_explicit_target_design_call() -> None:
    held = evaluate_architectural_delta(
        _request(
            _candidate(
                candidate_id="tonalide-default",
                material="Tonalide",
                family=ArchitecturalDeltaFamily.MUSK,
            )
        )
    )
    proposed = evaluate_architectural_delta(
        _request(
            _candidate(
                candidate_id="tonalide-exception",
                material="Tonalide",
                family=ArchitecturalDeltaFamily.MUSK,
                exception_justification=(
                    "The target explicitly requires a powdery polycyclic-musk shadow."
                ),
            )
        )
    )

    assert held.state is ArchitecturalDeltaState.HOLD
    assert any("exception-only" in blocker for blocker in held.blockers)
    assert proposed.state is ArchitecturalDeltaState.PROPOSED


def test_nary_design_requires_distinct_roles_and_complete_pairwise_isolates() -> None:
    held = evaluate_architectural_delta(
        _request(
            _candidate(
                candidate_id="layered-musks",
                material="Habanolide + Romandolide + Ambrettolide 10%",
                family=ArchitecturalDeltaFamily.MUSK,
                kind=ArchitecturalDeltaKind.NARY_DESIGN,
                distinct_role_refs=("radiance", "texture", "skin continuity"),
                pairwise_nonredundancy_refs=("habanolide:romandolide",),
            )
        )
    )
    contract_missing = evaluate_architectural_delta(
        _request(
            _candidate(
                candidate_id="layered-musks",
                material="Habanolide + Romandolide + Ambrettolide 10%",
                family=ArchitecturalDeltaFamily.MUSK,
                kind=ArchitecturalDeltaKind.NARY_DESIGN,
                distinct_role_refs=("radiance", "texture", "skin continuity"),
                pairwise_nonredundancy_refs=(
                    "habanolide:romandolide",
                    "habanolide:ambrettolide",
                    "romandolide:ambrettolide",
                ),
            )
        )
    )
    proposed = evaluate_architectural_delta(
        _request(
            _candidate(
                candidate_id="layered-musks",
                material="Habanolide + Romandolide + Ambrettolide 10%",
                family=ArchitecturalDeltaFamily.MUSK,
                kind=ArchitecturalDeltaKind.NARY_DESIGN,
                distinct_role_refs=("radiance", "texture", "skin continuity"),
                pairwise_nonredundancy_refs=(
                    "habanolide:romandolide",
                    "habanolide:ambrettolide",
                    "romandolide:ambrettolide",
                ),
                nary_candidate=_nary_contract(),
            )
        )
    )

    assert held.state is ArchitecturalDeltaState.HOLD
    assert any("pairwise" in blocker for blocker in held.blockers)
    assert contract_missing.state is ArchitecturalDeltaState.HOLD
    assert any("n-ary interaction contract" in blocker for blocker in contract_missing.blockers)
    assert proposed.state is ArchitecturalDeltaState.PROPOSED
    assert proposed.empirical_authority is False
